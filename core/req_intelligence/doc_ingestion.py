"""
Motor de ingesta de documentos para conversión a BDD.

Soporta:
  - Word (.docx): chunking semántico por encabezados H1/H2 y palabras clave en negrita.
  - Excel (.xlsx): chunking fila por fila con detección automática de columnas relevantes.

Produce objetos DocChunk que alimentan BDDDocConverter.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional


@dataclass
class DocChunk:
    """Unidad mínima de contenido extraído de un documento."""
    id: str
    source_file: str
    title: str
    content: str
    chunk_type: str  # "user_story" | "test_matrix"
    row_data: Optional[Dict[str, Any]] = field(default=None)


# ---------------------------------------------------------------------------
# Palabras clave que indican el inicio de una historia de usuario en Word
# ---------------------------------------------------------------------------
_STORY_KEYWORDS = re.compile(
    r"historia\s+de\s+usuario|user\s+story|criterios?\s+de\s+aceptaci[oó]n|"
    r"acceptance\s+criteria|caso\s+de\s+(prueba|uso)|test\s+case|escenario",
    re.IGNORECASE,
)

# Nombres de columnas que suelen contener contenido relevante en matrices Excel
_COLUMN_CANDIDATES = {
    "id": [
        r"id_caso", r"id caso", r"^id$", r"^id\b", "id del caso", "case id",
        "caso id", "num", "número", "n°",
    ],
    "name": ["nombre del caso", "nombre caso", "test case", "título del caso"],
    "description": [
        "descripci[oó]n", "description", "titulo", "título", "title",
        "historia", "user story",
    ],
    "preconditions": ["precondici[oó]n", "precondition", "prerrequisito"],
    "steps": [
        "pasos de ejecuci[oó]n", "pasos", "steps", "procedimiento", "procedure",
        "dado", "given", "acci[oó]n", "action",
    ],
    "expected": [
        "resultado esperado", "expected result", "resultado obtenido",
        "validaci[oó]n", "validation", "entonces", "then",
    ],
}

# Orden de asignación: roles más específicos primero; cada columna solo se usa una vez.
_COLUMN_ROLE_ORDER = ("id", "name", "description", "preconditions", "steps", "expected")

_SHEET_LOW_PRIORITY_NAME = re.compile(
    r"resumen|summary|sumario|índice|indice|portada|cover|instrucciones|legend|leyenda",
    re.IGNORECASE,
)

_TEST_ID_RE = re.compile(r"^[A-Z]{1,6}[-_]?\d+\b", re.IGNORECASE)


def _match_column(col_name: str, candidates: List[str]) -> bool:
    col_lower = col_name.lower().strip()
    col_norm = re.sub(r"[_]+", " ", col_lower)
    for pattern in candidates:
        if pattern.startswith("^") or pattern.endswith("$") or "\\b" in pattern:
            if re.search(pattern, col_norm, re.IGNORECASE):
                return True
        elif len(pattern) <= 4:
            if re.search(rf"\b{re.escape(pattern)}\b", col_norm, re.IGNORECASE):
                return True
        else:
            if re.search(pattern, col_lower, re.IGNORECASE):
                return True
    return False


def _detect_columns(df_columns: List[str]) -> Dict[str, Optional[str]]:
    """Detecta qué columna corresponde a cada rol (sin reutilizar la misma columna)."""
    mapping: Dict[str, Optional[str]] = {role: None for role in _COLUMN_ROLE_ORDER}
    used: set = set()
    for col in df_columns:
        col_s = str(col).strip()
        if not col_s or col_s.lower().startswith("unnamed"):
            continue
        for role in _COLUMN_ROLE_ORDER:
            if mapping[role] is not None:
                continue
            if col in used:
                break
            if _match_column(col_s, _COLUMN_CANDIDATES[role]):
                mapping[role] = col
                used.add(col)
                break
    return mapping


def _column_name_looks_like_header(name: str) -> bool:
    """True si el nombre de columna encaja con algún rol conocido."""
    for patterns in _COLUMN_CANDIDATES.values():
        if _match_column(name, patterns):
            return True
    return False


def _detect_columns_with_fallback(df_columns: List[str]) -> Dict[str, Optional[str]]:
    mapping = _detect_columns(df_columns)
    core = (mapping.get("id"), mapping.get("steps"), mapping.get("expected"))
    if any(core):
        return mapping
    cols = [c for c in df_columns if str(c).strip() not in ("", "nan")]
    if len(cols) < 2:
        return mapping
    return {
        "id": cols[0],
        "name": None,
        "description": cols[1] if len(cols) > 1 else None,
        "preconditions": cols[2] if len(cols) > 2 else None,
        "steps": cols[3] if len(cols) > 3 else None,
        "expected": cols[4] if len(cols) > 4 else None,
    }


def _prepare_excel_sheet(df_raw: Any) -> Any:
    """Usa la fila 0 como cabecera si ya parece matriz; si no, busca la fila de encabezados."""
    import pandas as pd  # type: ignore

    if df_raw is None or df_raw.empty:
        return df_raw
    header_hits = sum(1 for c in df_raw.columns if _column_name_looks_like_header(str(c)))
    if header_hits >= 2:
        df = df_raw.dropna(how="all").dropna(axis=1, how="all")
        return df if not df.empty else df_raw
    return _normalize_excel_sheet(df_raw)


def _score_test_matrix_sheet(df: Any, sheet_name: str) -> float:
    """Puntúa si la hoja parece una matriz de casos (no un resumen/tablas auxiliares)."""
    if df is None or df.empty:
        return 0.0
    col_map = _detect_columns_with_fallback(list(df.columns))
    score = 0.0
    if col_map.get("id"):
        score += 3.0
    if col_map.get("steps"):
        score += 3.0
    if col_map.get("expected"):
        score += 3.0
    if col_map.get("description") or col_map.get("name"):
        score += 2.0
    id_col = col_map.get("id")
    valid_ids = 0
    if id_col:
        for _, row in df.iterrows():
            if _TEST_ID_RE.match(_cell_text(row[id_col])):
                valid_ids += 1
    score += min(valid_ids, 80) * 0.35
    n_rows = len(df)
    if n_rows >= 10:
        score += 2.0
    score += min(n_rows, 120) * 0.05
    if _SHEET_LOW_PRIORITY_NAME.search(sheet_name or ""):
        score *= 0.25
    if n_rows <= 10 and len(df.columns) <= 4:
        score *= 0.35
    return score


def _normalize_excel_sheet(df_raw: Any) -> Any:
    """Elimina filas/columnas vacías y, si hace falta, localiza la fila de encabezados."""
    import pandas as pd  # type: ignore

    if df_raw is None or df_raw.empty:
        return df_raw
    df = df_raw.dropna(how="all").dropna(axis=1, how="all")
    if df.empty:
        return df

    cols_str = [str(c).strip() for c in df.columns]
    unnamed = sum(1 for c in cols_str if c.startswith("Unnamed") or c.isdigit())
    if unnamed < len(cols_str) and any(_column_name_looks_like_header(c) for c in cols_str):
        return df

    max_scan = min(20, len(df))
    for i in range(max_scan):
        row_vals = [str(v).strip() for v in df.iloc[i].tolist()]
        named = [v for v in row_vals if v and v.lower() not in ("nan", "none")]
        if len(named) < 2:
            continue
        header_hits = sum(1 for v in named if _column_name_looks_like_header(v))
        if header_hits >= 1 or (len(named) >= 3 and i > 0):
            new_df = df.iloc[i + 1 :].copy()
            new_df.columns = row_vals[: len(new_df.columns)]
            new_df = new_df.dropna(how="all").dropna(axis=1, how="all")
            if not new_df.empty:
                return new_df
    return df


def _read_excel_all_sheets(file_path: str, pd: Any) -> Dict[str, Any]:
    """Lee todas las hojas; prueba calamine, openpyxl y motor por defecto."""
    last_err: Optional[Exception] = None
    for engine in ("calamine", "openpyxl", None):
        try:
            kwargs: Dict[str, Any] = {"sheet_name": None}
            if engine:
                kwargs["engine"] = engine
            sheets = pd.read_excel(file_path, **kwargs)
            if isinstance(sheets, dict):
                return sheets
            return {"": sheets}
        except Exception as e:
            last_err = e
    if last_err:
        raise last_err
    return {}


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    if s.lower() in ("nan", "none", ""):
        return ""
    if "dtype:" in s:
        return ""
    return s


def _chunks_from_dataframe(
    df: Any,
    file_path: str,
    basename: str,
    sheet_label: str,
) -> List[DocChunk]:
    """Convierte filas de una hoja en DocChunk."""
    df = _prepare_excel_sheet(df)
    if df is None or df.empty:
        return []

    col_map = _detect_columns_with_fallback(list(df.columns))
    chunks: List[DocChunk] = []
    sheet_suffix = f"_{sheet_label}" if sheet_label else ""
    used_cols = {c for c in col_map.values() if c}

    for index, row in df.iterrows():
        row_id = _cell_text(row[col_map["id"]]) if col_map.get("id") else ""
        case_name = _cell_text(row[col_map["name"]]) if col_map.get("name") else ""
        description = _cell_text(row[col_map["description"]]) if col_map.get("description") else ""
        preconditions = _cell_text(row[col_map["preconditions"]]) if col_map.get("preconditions") else ""
        steps = _cell_text(row[col_map["steps"]]) if col_map.get("steps") else ""
        expected = _cell_text(row[col_map["expected"]]) if col_map.get("expected") else ""

        if row_id.lower() in ("id_caso", "id caso", "id", "caso", "n°", "num"):
            continue
        if not row_id:
            row_id = str(index + 1)
        elif not _TEST_ID_RE.match(row_id) and not any([case_name, description, steps, expected]):
            continue

        if not any([case_name, description, steps, expected]):
            extras: List[str] = []
            for col in df.columns:
                if col in used_cols:
                    continue
                txt = _cell_text(row[col])
                if txt:
                    extras.append(f"{col}: {txt}")
            if extras:
                description = " | ".join(extras[:4])
            else:
                continue

        content_parts: List[str] = []
        if case_name:
            content_parts.append(f"Nombre: {case_name}")
        if description:
            content_parts.append(f"Descripción: {description}")
        if preconditions:
            content_parts.append(f"Precondiciones: {preconditions}")
        if steps:
            content_parts.append(f"Pasos: {steps}")
        if expected:
            content_parts.append(f"Resultado esperado: {expected}")

        title = f"Caso {row_id}"
        if case_name:
            title += f": {case_name[:70]}"
        elif description:
            title += f": {description[:70]}"

        chunks.append(DocChunk(
            id=f"{basename}{sheet_suffix}_{index}",
            source_file=file_path,
            title=title,
            content="\n".join(content_parts),
            chunk_type="test_matrix",
            row_data={
                "id": row_id,
                "name": case_name,
                "description": description,
                "preconditions": preconditions,
                "steps": steps,
                "expected": expected,
                "sheet": sheet_label,
            },
        ))

    return chunks


# ---------------------------------------------------------------------------
# WordIngester
# ---------------------------------------------------------------------------

class WordIngester:
    """
    Extrae chunks de un documento Word (.docx).

    Estrategia:
    1. Itera párrafos del documento.
    2. Cuando encuentra un encabezado (H1/H2) o un párrafo en negrita que coincide
       con palabras clave de historias, inicia un nuevo chunk.
    3. Acumula el texto hasta el siguiente encabezado clave.
    4. Cada chunk representa una historia de usuario o criterio de aceptación.
    """

    def ingest(self, file_path: str) -> List[DocChunk]:
        try:
            from docx import Document  # type: ignore
        except ImportError as e:
            raise ImportError(
                "python-docx no está instalado. Ejecuta: pip install python-docx"
            ) from e

        doc = Document(file_path)
        basename = os.path.basename(file_path)
        chunks: List[DocChunk] = []
        current_title = ""
        current_lines: List[str] = []
        chunk_index = 0

        def _flush() -> None:
            nonlocal chunk_index
            content = "\n".join(current_lines).strip()
            if content and current_title:
                chunks.append(DocChunk(
                    id=f"{basename}_{chunk_index}",
                    source_file=file_path,
                    title=current_title,
                    content=content,
                    chunk_type="user_story",
                ))
                chunk_index += 1

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            is_heading = para.style.name.startswith("Heading") and int(
                para.style.name.split()[-1]
            ) <= 2 if para.style.name.startswith("Heading") else False
            is_bold_keyword = (
                any(run.bold for run in para.runs if run.text.strip())
                and _STORY_KEYWORDS.search(text)
            )

            if is_heading or is_bold_keyword:
                _flush()
                current_title = text
                current_lines = []
            else:
                current_lines.append(text)

        _flush()

        # Si no se detectaron chunks con headings, tratar todo el documento como un único chunk
        if not chunks:
            full_text = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
            if full_text:
                chunks.append(DocChunk(
                    id=f"{basename}_0",
                    source_file=file_path,
                    title=basename,
                    content=full_text,
                    chunk_type="user_story",
                ))

        return chunks


# ---------------------------------------------------------------------------
# ExcelIngester
# ---------------------------------------------------------------------------

class ExcelIngester:
    """
    Extrae chunks de un archivo Excel (.xlsx).

    Estrategia:
    1. Lee todas las hojas (la matriz suele estar en la 2ª hoja o posteriores).
    2. Normaliza encabezados si la primera fila no es la cabecera real.
    3. Detecta columnas por nombre o, en su defecto, por posición.
    4. Genera un DocChunk por cada fila con contenido útil.
    """

    def ingest(self, file_path: str) -> List[DocChunk]:
        try:
            import pandas as pd  # type: ignore
        except ImportError as e:
            raise ImportError(
                "pandas no está instalado. Desde el venv del proyecto ejecuta: "
                "pip install -r requirements.txt "
                "(necesita pandas, python-calamine y openpyxl). "
                "Si usas ELIA.exe, vuelve a generar el ejecutable tras actualizar ELIA.spec."
            ) from e

        basename = os.path.basename(file_path)
        sheets = _read_excel_all_sheets(file_path, pd)
        if not sheets:
            return []

        scored: List[tuple] = []
        for sheet_name, df_raw in sheets.items():
            label = str(sheet_name).strip() if sheet_name is not None else ""
            prepared = _prepare_excel_sheet(df_raw)
            score = _score_test_matrix_sheet(prepared, label)
            scored.append((label, score, df_raw))

        scored.sort(key=lambda item: item[1], reverse=True)
        best_score = scored[0][1] if scored else 0.0
        threshold = max(8.0, best_score * 0.55)

        chunks: List[DocChunk] = []
        for label, score, df_raw in scored:
            if score < threshold:
                continue
            chunks.extend(_chunks_from_dataframe(df_raw, file_path, basename, label))

        if not chunks and scored:
            label, _, df_raw = scored[0]
            chunks = _chunks_from_dataframe(df_raw, file_path, basename, label)

        return chunks
